import datetime
import os
import platform
from pathlib import Path

import psutil

MAX_LIST_ENTRIES = 80
MAX_READ_CHARS = 80_000
DOCUMENT_EXTENSIONS = {
    ".docx": "Word",
    ".pdf": "PDF",
    ".xlsx": "Excel",
    ".xlsm": "Excel",
    ".pptx": "PowerPoint",
}


def get_time():
    now = datetime.datetime.now().strftime("%I:%M %p")
    return f"The current time is {now}."


def get_system_status():
    cpu = psutil.cpu_percent(interval=0.2)
    ram = psutil.virtual_memory().percent
    if platform.system() == "Windows":
        root = os.path.splitdrive(os.getcwd())[0] + "\\"
    else:
        root = "/"
    disk = psutil.disk_usage(root).percent
    return (
        f"ZEUS System Health: OS: {platform.system()} | "
        f"CPU: {cpu}% | RAM: {ram}% | Disk: {disk}%"
    )


def list_files(directory="."):
    try:
        path = (directory or ".").strip() or "."
        names = os.listdir(path)
        extra = ""
        if len(names) > MAX_LIST_ENTRIES:
            extra = f" ... (+{len(names) - MAX_LIST_ENTRIES} more)"
            names = names[:MAX_LIST_ENTRIES]
        return f"Files in '{path}': " + ", ".join(names) + extra
    except Exception as e:
        return f"Error accessing path: {e}"


def read_file(filepath):
    try:
        path = Path(filepath.strip()).expanduser()
        if not path.is_file():
            return f"Error reading file: '{path}' is not a file."

        if path.suffix.lower() in DOCUMENT_EXTENSIONS:
            text = _read_document(path)
        else:
            with path.open("r", encoding="utf-8", errors="replace") as file_handle:
                text = file_handle.read(MAX_READ_CHARS + 1)
        if len(text) > MAX_READ_CHARS:
            text = text[:MAX_READ_CHARS] + "\n... [truncated]"
        return f"Content of {path}:\n{text}"
    except Exception as e:
        return f"Error reading file: {e}"


def _read_document(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".docx":
        from docx import Document

        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)

    if extension == ".pdf":
        from pypdf import PdfReader

        pages = []
        for page in PdfReader(str(path)).pages:
            pages.append(page.extract_text() or "")
        return "\n\n".join(pages)

    if extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = []
            for worksheet in workbook.worksheets:
                rows = []
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values):
                        rows.append(" | ".join(values))
                sheets.append(f"[Sheet: {worksheet.title}]\n" + "\n".join(rows))
            return "\n\n".join(sheets)
        finally:
            workbook.close()

    if extension == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        slides = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            slides.append(f"[Slide {index}]\n" + "\n".join(texts))
        return "\n\n".join(slides)

    raise ValueError(f"Unsupported document format: {extension}")
