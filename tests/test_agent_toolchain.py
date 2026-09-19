from agent import ZeusAgent
from tools.system_tools import read_file


def test_extract_tool_names_handles_chained_tools():
    agent = ZeusAgent.__new__(ZeusAgent)
    tags = agent._extract_tool_names('[[GET_TIME]] and [[SEARCH_WEB: AI trends]]')
    assert tags == ['GET_TIME', 'SEARCH_WEB']


def test_markdown_table_is_rendered_as_html_table():
    text = '| Name | Status |\n| --- | --- |\n| Zeus | Online |'
    html = ZeusAgent.render_markdown_table(text)
    assert '<table' in html.lower()
    assert 'electric-cyan' in html.lower()


def test_read_file_extracts_word_excel_and_powerpoint(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    word_path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("Word content")
    document.save(word_path)

    excel_path = tmp_path / "data.xlsx"
    workbook = Workbook()
    workbook.active.append(["Name", "Status"])
    workbook.active.append(["ZEUS", "Online"])
    workbook.save(excel_path)

    powerpoint_path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Presentation content"
    presentation.save(powerpoint_path)

    assert "Word content" in read_file(str(word_path))
    assert "ZEUS | Online" in read_file(str(excel_path))
    assert "Presentation content" in read_file(str(powerpoint_path))
